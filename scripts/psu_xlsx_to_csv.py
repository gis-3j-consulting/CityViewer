#!/usr/bin/env python3
"""Convert PSU Population Research Center's certified estimates workbook into data/psu_population.csv.

    python3 scripts/psu_xlsx_to_csv.py path/to/2025_Certified_Population_Estimates_.xlsx

Reads the "Cities and Towns" sheet, which PSU lays out as several side-by-side blocks of
(city name, April 1 estimates base, then July 1 estimates for each year). Keeps the July 1
columns, strips footnote asterisks and rounds to whole people. Needs: pip install openpyxl
"""
import csv
import sys
from pathlib import Path

import openpyxl

src = Path(sys.argv[1])
dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / "data" / "psu_population.csv"

ws = openpyxl.load_workbook(src, data_only=True)["Cities and Towns"]

# Find each block: a header cell reading "Incorporated City/Town" on the first header rows.
header_row = next(r for r in range(1, 8) if any(ws.cell(row=r, column=c).value == "Incorporated City/Town" for c in range(1, ws.max_column + 1)))
date_row = header_row + 1
name_cols = [c for c in range(1, ws.max_column + 1) if ws.cell(row=header_row, column=c).value == "Incorporated City/Town"]

rows, years = [], None
for c0 in name_cols:
    dates = [ws.cell(row=date_row, column=c0 + k).value for k in range(1, 8)]
    ys = [d.year for d in dates]
    # first date is the April 1 estimates base; the other six are July 1 of consecutive years
    assert len(ys) == 7 and ys[1:] == list(range(ys[1], ys[1] + 6)), f"Unexpected date headers: {dates}"
    years = ys[1:]
    for r in range(date_row + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=c0).value
        vals = [ws.cell(row=r, column=c0 + k).value for k in range(2, 8)]
        if name is None or not all(isinstance(v, (int, float)) for v in vals):
            continue
        rows.append([str(name).replace("*", "").strip()] + [int(round(v)) for v in vals])

rows.sort(key=lambda r: r[0].lower())
with open(dst, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["city"] + years)
    w.writerows(rows)
print(f"Wrote {len(rows)} cities ({years[0]}-{years[-1]}) to {dst}")
